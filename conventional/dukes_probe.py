#!/usr/bin/env python3
"""Temporary issue #76 probe: inspect 2026 DUKES 5.11 headers and identified evidence rows."""
from __future__ import annotations
import hashlib
import io
import sys
import urllib.request
from openpyxl import load_workbook
URL = "https://assets.publishing.service.gov.uk/media/6a6a365b0ddb7e4831c629f3/DUKES_5.11.xlsx"
with urllib.request.urlopen(URL, timeout=60) as response:
    data = response.read(2_000_000)
digest = hashlib.sha256(data).hexdigest()
print("DUKES bytes", len(data), "sha256", digest)
if len(data) != 1247443 or not digest.startswith("760286f4cb3168f0"):
    raise SystemExit("DUKES source changed: refuse to inspect")
wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
print("SHEETS:", wb.sheetnames)
for ws in wb:
    print("===SHEET", repr(ws.title), "rows", ws.max_row, "cols", ws.max_column)
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        vals = [(j + 1, str(v).strip()[:85]) for j, v in enumerate(row) if v is not None and str(v).strip()]
        if i <= 14 and vals:
            print("HEADER", i, vals[:24])
        if any(k in " ".join(str(v) for v in row if v is not None).upper() for k in ("MARCHWOOD", "KILROOT", "SIZEWELL B", "DRAX")):
            print("MATCH", i, vals[:24])

from collections import Counter
from conventional.fetchers.dukes_5_11 import normalise, RAW_FILE
stations = normalise({RAW_FILE: data})
print("DUKES_SOURCE_ROWS", len(stations), "BY_FUEL", dict(sorted(Counter(v["fuel"] for v in stations).items())))
print("DUKES_MARCHWOOD", [(x["id"], x["name"], x["capacity_mw"], x["fuel"]) for x in stations if "marchwood" in x["name"].lower()])

import json
from pathlib import Path
from conventional.dukes_reconcile import reconcile, encode
baseline = json.loads((Path(__file__).resolve().parent / "lanes/dc-a6/out/conventional_power_uk.geojson").read_text())
derived, audit = reconcile(baseline, stations)
print("DUKES_JOIN_COUNTS", {k:v for k,v in audit.items() if not isinstance(v,list)})
print("DUKES_MATCHES", audit["matches"][:80])
print("DUKES_REJECTED", audit["rejected_matches"][:50])
print("DUKES_OUTPUT_SHA256", hashlib.sha256(encode(derived)).hexdigest())
