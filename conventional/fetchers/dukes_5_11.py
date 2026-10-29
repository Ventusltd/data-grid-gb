"""DUKES Table 5.11 (DESNZ): deterministic, pinned May 2026 station records.

Source workbook inspected on GitHub Actions run 40961887500 (requirements leg).
Sheet "5.11 Full list", header on row 6: site, technology, fuel, installed MW,
postcode, OS coordinates, year, unique DESNZ site code. Operational MPP sites
only. These rows DO NOT establish connection rights, DNC, export MW, or capacity
at any separately mapped GridAtlas station.

No geocoding or guesses: DUKES supplies British National Grid coordinates,
not WGS84. Keep the raw grid references and return a null GeoJSON geometry.
Only a separate, unique, source-evidenced site reconciliation may link them
to the map; never use name similarity or physical proximity as identity.
"""
from __future__ import annotations

import math
import re
from io import BytesIO

SOURCE_ID = "dukes-5-11"
LANE = "DC-A6"
STATUS = "ready"
URL = "https://assets.publishing.service.gov.uk/media/6a6a365b0ddb7e4831c629f3/DUKES_5.11.xlsx"
RAW_FILE = "DUKES_5.11-2026-05.xlsx"
SOURCE_SHA256 = "760286f4cb3168f04a5ee1b581dbfcb35e1d65a5ea0151bb2dd4dc12dca2c72a"
DATE = "2026-05"
PUBLISHED = "2026-07-30"
SHEET = "5.11 Full list"
EXPECTED_COLUMNS = (
    "Company Name [note 30]", "Site Name", "Technology", "Type", "CHP",
    "Primary Fuel", "Secondary Fuel", "InstalledCapacity (MW)",
    "Grid Connection Type", "Country", "Region", "Postcode",
    "OS Reference", "X-Coordinate", "Y-Coordinate", "Year Commissioned",
    "DESNZ site code",
)
FUEL = {
    "natural gas": "gas", "sour gas": "gas", "coal": "coal",
    "oil": "oil", "diesel/gas oil": "gas_oil_dual",
    "diesel": "diesel", "nuclear": "nuclear",
    "biomass": "biomass", "waste": "efw",
}


def raw_requests():
    return [{"file": RAW_FILE, "url": URL}]


def _text(v):
    return re.sub(r"\s+", " ", str(v or "")).strip()


def normalise(raw):
    """Return DUKES station rows; duplicates or altered sheet layouts fail closed.

    This is deliberately NOT a fuzzy matcher or a converter of OSGB to WGS84.
    Unmatched rows and unknown MW stay explicit. Only installed MW is taken.
    """
    import hashlib
    payload = raw[RAW_FILE]
    got = hashlib.sha256(payload).hexdigest()
    if got != SOURCE_SHA256:
        raise ValueError("DUKES May 2026 workbook sha256 differs: " + got)
    try:
        from openpyxl import load_workbook
    except ImportError as ex:
        raise RuntimeError("DUKES XLSX normalisation requires openpyxl") from ex
    book = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    if SHEET not in book.sheetnames:
        raise ValueError("DUKES sheet 5.11 Full list absent")
    sheet = book[SHEET]
    rows = sheet.iter_rows(values_only=True)
    first = next(rows)
    if "operational" not in _text(first[0]).lower():
        raise ValueError("DUKES May 2026 operational-only scope not confirmed")
    for _ in range(4):
        next(rows)
    header = tuple(_text(v) for v in next(rows))
    if header != EXPECTED_COLUMNS:
        raise ValueError("DUKES header changed: " + repr(header))
    output, seen = [], set()
    for excel_row, row in enumerate(rows, start=7):
        if len(row) != 17:
            raise ValueError(f"DUKES row {excel_row}: expected 17 columns")
        company, site, technology, unit_type, chp, fuel, secondary, mw, connection, country, region, postcode, osref, east, north, year, sitecode = row
        fuel_raw = _text(fuel)
        kind = FUEL.get(fuel_raw.casefold())
        if kind is None:
            continue
        name = _text(site)
        code = _text(sitecode)
        if not name or not re.fullmatch(r"[A-Z]{3}-[0-9]{6}", code):
            # No unique source identity: not safe to put into the join universe.
            continue
        identity = "dukes-5-11:" + code.lower()
        if identity in seen:
            raise ValueError(f"DUKES duplicate source site code {code}; do not sum")
        seen.add(identity)
        cap = None
        if mw is not None and _text(mw):
            try:
                cap = float(mw)
            except (TypeError, ValueError) as ex:
                raise ValueError(f"DUKES {code}: invalid installed MW {mw!r}") from ex
            if not math.isfinite(cap) or cap <= 0:
                raise ValueError(f"DUKES {code}: zero/negative/invalid MW, unknown cannot be zero")
        commissioned = None
        if year is not None and _text(year):
            try:
                candidate = int(year)
                if 1880 <= candidate <= 2100:
                    commissioned = candidate
            except (TypeError, ValueError):
                pass
        output.append({
            "id": identity,
            "name": name,
            "operator": _text(company) or None,
            "fuel": kind,
            "primary_fuel_as_published": fuel_raw,
            "secondary_fuel_as_published": _text(secondary) or None,
            "technology": _text(technology) or None,
            "technology_type": _text(unit_type) or None,
            "chp_as_published": _text(chp) or None,
            "grid_connection_type_as_published": _text(connection) or None,
            "postcode_as_published": _text(postcode) or None,
            "os_reference_as_published": _text(osref) or None,
            "osgb_easting_m_as_published": east,
            "osgb_northing_m_as_published": north,
            "desnz_site_code": code,
            "source_excel_row": excel_row,
            "source_data_date": DATE,
            "source_publication_date": PUBLISHED,
            "status": "operational",
            "status_basis": 'DUKES 5.11 "Full list": MPP station operational as at May 2026',
            "capacity_mw": cap,
            "capacity_basis": 'DUKES 5.11 "5.11 Full list", InstalledCapacity (MW), May 2026; NOT export, TEC, DNC or grid-connection MW',
            "commissioned_year": commissioned,
            "location": {"precision": "unknown", "lon": None, "lat": None},
            "source": {
                "id": SOURCE_ID, "licence": "OGL-3.0",
                "licence_url": "https://www.nationalarchives.gov.uk/doc/open-government-licence/version/3/",
                "evidence_url": URL,
            },
        })
    return sorted(output, key=lambda rec: rec["id"])
