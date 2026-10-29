"""Read a worksheet from raw .xlsx bytes (DUKES 5.11 and similar government tables).

openpyxl is the only optional dependency of the asset pipelines (requirements.txt).
Fetching, pinning, verifying and reproducing never need it; only a normaliser that
reads a spreadsheet does, and it fails with a clear message when it is absent.
"""
from __future__ import annotations

import io


class MissingDependency(RuntimeError):
    pass


def available() -> bool:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return False
    return True


def sheet_rows(data: bytes, sheet: str) -> list:
    """Every row of ``sheet`` as a list of plain values (str, int, float, None)."""
    try:
        import openpyxl
    except ImportError as error:
        raise MissingDependency("reading .xlsx sources needs openpyxl: pip install -r requirements.txt") from error
    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        if sheet not in book.sheetnames:
            raise KeyError(f"sheet {sheet!r} not in {book.sheetnames}")
        return [list(row) for row in book[sheet].iter_rows(values_only=True)]
    finally:
        book.close()
